"""Download the original UCI workbook and stream it into a Spark-readable CSV."""
import argparse
import csv
import shutil
import urllib.request
import zipfile
from pathlib import Path

UCI_URL = "https://archive.ics.uci.edu/static/public/352/online+retail.zip"
HEADERS = ["InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate",
           "UnitPrice", "CustomerID", "Country"]


def convert_workbook(source: Path, destination: Path) -> int:
    from openpyxl import load_workbook

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(source, read_only=True, data_only=True)
    count = 0
    temporary = destination.with_suffix(".csv.tmp")
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = list(next(rows))
        indexes = [header.index(name) for name in HEADERS]
        with temporary.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(HEADERS)
            for row in rows:
                values = [row[i] for i in indexes]
                if hasattr(values[4], "strftime"):
                    values[4] = values[4].strftime("%Y-%m-%d %H:%M:%S")
                writer.writerow(values)
                count += 1
        temporary.replace(destination)
    finally:
        workbook.close()
    return count


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--directory", type=Path, default=Path("data/raw"))
    parser.add_argument("--xlsx", type=Path, help="Use a previously downloaded workbook")
    args = parser.parse_args()
    args.directory.mkdir(parents=True, exist_ok=True)
    workbook = args.xlsx
    if workbook is None:
        archive = args.directory / "online-retail.zip"
        if not archive.exists():
            temporary = archive.with_suffix(".zip.tmp")
            print("Downloading UCI Online Retail (about 23 MB)...")
            with urllib.request.urlopen(UCI_URL, timeout=120) as response:
                with temporary.open("wb") as handle:
                    shutil.copyfileobj(response, handle)
            temporary.replace(archive)
        workbook = args.directory / "Online Retail.xlsx"
        with zipfile.ZipFile(archive) as bundle:
            member = next(name for name in bundle.namelist() if name.endswith(".xlsx"))
            with bundle.open(member) as source, workbook.open("wb") as target:
                shutil.copyfileobj(source, target)
    output = args.directory / "online_retail.csv"
    count = convert_workbook(workbook, output)
    print(f"Prepared {count:,} transaction rows: {output}")


if __name__ == "__main__":
    main()
